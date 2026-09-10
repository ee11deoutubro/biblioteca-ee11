from __future__ import annotations

import re
import unicodedata
from collections import Counter
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
WORKBOOK = ROOT.parent / "upload" / "RELAÇÃO GERAL DE TÍTULOS E AUTORES(2).xlsx"
OUTPUT = ROOT / "importar-metadados-planilha-original.sql"

CATEGORIES = {
    "LBC": "Literatura brasileira contemporânea",
    "LE": "Literatura estrangeira",
    "LJ": "Literatura juvenil",
    "CONT": "Contos",
}

COLOR_HEX = {
    "VERDE": "#16a34a",
    "AMARELA": "#eab308",
    "ROSA": "#ec4899",
    "LARANJA": "#f59e0b",
}

# Linhas usadas apenas para completar modelos vazios da planilha.
PLACEHOLDERS = {
    "nao sei o que la",
    "anarcopolis",
    "ancapsu",
}


def normalize(value: object) -> str:
    text = unicodedata.normalize("NFD", str(value or "").strip().lower())
    text = "".join(char for char in text if unicodedata.category(char) != "Mn")
    return re.sub(r"\s+", " ", text).strip()


def clean(value: object) -> str | None:
    text = str(value or "").strip()
    return text or None


def sql_text(value: object) -> str:
    if value is None:
        return "null"
    return "'" + str(value).replace("'", "''") + "'"


workbook = load_workbook(WORKBOOK, data_only=True, read_only=True)

colors_by_title: dict[str, Counter[str]] = {}
planilha_titles: dict[str, tuple[str, str | None]] = {}
for row in workbook["Planilha1"].iter_rows(min_row=2, values_only=True):
    title = clean(row[0] if len(row) > 0 else None)
    author = clean(row[1] if len(row) > 1 else None)
    color = clean(row[2] if len(row) > 2 else None)
    key = normalize(title)
    if not key or key in PLACEHOLDERS:
        continue
    planilha_titles.setdefault(key, (title, author))
    if color:
        colors_by_title.setdefault(key, Counter())[color.upper()] += 1

# A ordem das abas também resolve os dois títulos repetidos entre categorias:
# LBC prevalece sobre LJ em "Assassinato na Biblioteca" e LE prevalece sobre
# LJ em "Confissões de um Amigo Imaginário", reproduzindo os 398 títulos atuais.
records: dict[str, dict[str, object]] = {}
for code, category in CATEGORIES.items():
    for row in workbook[code].iter_rows(min_row=2, values_only=True):
        order = row[0] if len(row) > 0 else None
        title = clean(row[1] if len(row) > 1 else None)
        author = clean(row[2] if len(row) > 2 else None)
        location = clean(row[4] if len(row) > 4 else None)
        key = normalize(title)
        if not key or key in PLACEHOLDERS or key in records:
            continue

        records[key] = {
            "titulo_chave": key,
            "titulo_original": title,
            "autor_original": author,
            "categoria": category,
            "ordem_planilha": int(order) if isinstance(order, (int, float)) else None,
            "genero_codigo": code,
            "classificacao_numero": location,
        }

# A aba de cores contém 27 títulos juvenis que não aparecem nas quatro abas
# principais. Eles pertencem ao catálogo atual, mas não possuem ordem/localização.
for key, (title, author) in planilha_titles.items():
    if key in records:
        continue
    records[key] = {
        "titulo_chave": key,
        "titulo_original": title,
        "autor_original": author,
        "categoria": CATEGORIES["LJ"],
        "ordem_planilha": None,
        "genero_codigo": "LJ",
        "classificacao_numero": None,
    }

for key, record in records.items():
    color_counter = colors_by_title.get(key, Counter())
    source_color = color_counter.most_common(1)[0][0] if color_counter else None
    record["classificacao_cor"] = source_color.title() if source_color else None
    record["classificacao_cor_hex"] = COLOR_HEX.get(source_color)

if len(records) != 398:
    raise RuntimeError(f"A conferência deveria resultar em 398 títulos, mas encontrou {len(records)}.")

values = []
for record in records.values():
    values.append(
        "(" + ",".join([
            sql_text(record["titulo_chave"]),
            sql_text(record["titulo_original"]),
            sql_text(record["autor_original"]),
            sql_text(record["categoria"]),
            str(record["ordem_planilha"]) if record["ordem_planilha"] is not None else "null",
            sql_text(record["genero_codigo"]),
            sql_text(record["classificacao_numero"]),
            sql_text(record["classificacao_cor"]),
            sql_text(record["classificacao_cor_hex"]),
        ]) + ")"
    )

sql = f"""-- Biblioteca EE 11 de Outubro
-- Importa somente metadados da planilha original nos 398 títulos existentes.
-- Não cria títulos, não altera quantidades e não apaga exemplares.

begin;

alter table public.livros
  add column if not exists ordem_planilha integer,
  add column if not exists genero_codigo text,
  add column if not exists classificacao_numero text,
  add column if not exists classificacao_cor text,
  add column if not exists classificacao_cor_hex text;

alter table public.exemplares
  add column if not exists tombamento text;

create or replace function pg_temp.normalizar_catalogo(valor text)
returns text
language sql
immutable
as $$
  select btrim(
    regexp_replace(
      translate(
        lower(coalesce(valor, '')),
        'áàãâäéèêëíìîïóòõôöúùûüçñ',
        'aaaaaeeeeiiiiooooouuuucn'
      ),
      '\\s+',
      ' ',
      'g'
    )
  );
$$;

create temporary table catalogo_planilha_original (
  titulo_chave text primary key,
  titulo_original text not null,
  autor_original text,
  categoria text not null,
  ordem_planilha integer,
  genero_codigo text,
  classificacao_numero text,
  classificacao_cor text,
  classificacao_cor_hex text
) on commit preserve rows;

insert into catalogo_planilha_original (
  titulo_chave,
  titulo_original,
  autor_original,
  categoria,
  ordem_planilha,
  genero_codigo,
  classificacao_numero,
  classificacao_cor,
  classificacao_cor_hex
) values
{',\n'.join(values)};

update public.livros as livro
set
  ordem_planilha = coalesce(livro.ordem_planilha, origem.ordem_planilha),
  genero_codigo = coalesce(nullif(btrim(livro.genero_codigo), ''), origem.genero_codigo),
  classificacao_numero = coalesce(nullif(btrim(livro.classificacao_numero), ''), origem.classificacao_numero),
  classificacao_cor = coalesce(nullif(btrim(livro.classificacao_cor), ''), origem.classificacao_cor),
  classificacao_cor_hex = coalesce(nullif(btrim(livro.classificacao_cor_hex), ''), origem.classificacao_cor_hex)
from catalogo_planilha_original as origem
where pg_temp.normalizar_catalogo(livro.titulo) = origem.titulo_chave;

update public.exemplares as exemplar
set localizacao = origem.classificacao_numero
from public.livros as livro
join catalogo_planilha_original as origem
  on pg_temp.normalizar_catalogo(livro.titulo) = origem.titulo_chave
where exemplar.livro_id = livro.id
  and nullif(btrim(exemplar.localizacao), '') is null
  and origem.classificacao_numero is not null;

commit;

select
  (select count(*) from catalogo_planilha_original) as titulos_na_planilha,
  count(distinct livro.id) as titulos_localizados_no_banco,
  count(distinct livro.id) filter (where livro.ordem_planilha is not null) as com_ordem,
  count(distinct livro.id) filter (where livro.genero_codigo is not null) as com_codigo,
  count(distinct livro.id) filter (where livro.classificacao_numero is not null) as com_numero_ou_localizacao,
  count(distinct livro.id) filter (where livro.classificacao_cor is not null) as com_cor_original,
  coalesce(
    jsonb_agg(
      jsonb_build_object(
        'titulo', origem.titulo_original,
        'autor', origem.autor_original
      ) order by origem.titulo_original
    ) filter (where livro.id is null),
    '[]'::jsonb
  ) as titulos_nao_localizados
from catalogo_planilha_original as origem
left join public.livros as livro
  on pg_temp.normalizar_catalogo(livro.titulo) = origem.titulo_chave;
"""

OUTPUT.write_text(sql, encoding="utf-8")
print(f"{len(records)} títulos gravados em {OUTPUT.name}")
