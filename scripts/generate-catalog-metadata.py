from __future__ import annotations

import json
import re
import unicodedata
from collections import Counter
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
WORKBOOK = ROOT.parent / "upload" / "RELAÇÃO GERAL DE TÍTULOS E AUTORES(2).xlsx"
OUTPUT = ROOT / "catalogo-metadata-original.js"

CATEGORIES = {
    "LBC": "Literatura brasileira contemporânea",
    "LE": "Literatura estrangeira",
    "LJ": "Literatura juvenil",
    "CONT": "Contos",
}

CATEGORY_COLORS = {
    "LBC": "#f59e0b",
    "LE": "#7c3aed",
    "LJ": "#ec4899",
    "CONT": "#2563eb",
}

COLOR_HEX = {
    "AZUL": "#2563eb",
    "AMARELO": "#eab308",
    "LARANJA": "#f59e0b",
    "ROSA": "#ec4899",
    "VERDE": "#16a34a",
    "VERMELHO": "#dc2626",
    "ROXO": "#7c3aed",
    "MARROM": "#92400e",
    "PRETO": "#111827",
    "BRANCO": "#cbd5e1",
}


def normalize(value: object) -> str:
    text = unicodedata.normalize("NFD", str(value or "").strip().lower())
    text = "".join(char for char in text if unicodedata.category(char) != "Mn")
    return re.sub(r"\s+", " ", text)


def clean(value: object) -> str | None:
    text = str(value or "").strip()
    return text or None


book = load_workbook(WORKBOOK, data_only=True, read_only=True)

colors_by_title: dict[str, Counter[str]] = {}
for row in book["Planilha1"].iter_rows(min_row=2, values_only=True):
    title = normalize(row[0] if len(row) > 0 else None)
    color = clean(row[2] if len(row) > 2 else None)
    if title and color:
        colors_by_title.setdefault(title, Counter())[color.upper()] += 1

metadata: dict[str, dict[str, object]] = {}
for code, category in CATEGORIES.items():
    for row in book[code].iter_rows(min_row=2, values_only=True):
        order = row[0] if len(row) > 0 else None
        title = clean(row[1] if len(row) > 1 else None)
        location = clean(row[4] if len(row) > 4 else None)
        key = normalize(title)
        if not key or key in metadata:
            continue

        color_counter = colors_by_title.get(key, Counter())
        source_color = color_counter.most_common(1)[0][0] if color_counter else None
        color_label = source_color.title() if source_color else None
        metadata[key] = {
            "ordem_planilha": int(order) if isinstance(order, (int, float)) else None,
            "categoria": category,
            "genero_codigo": code,
            "classificacao_numero": location,
            "classificacao_cor": color_label,
            "classificacao_cor_hex": COLOR_HEX.get(source_color, CATEGORY_COLORS[code]),
        }

payload = json.dumps(metadata, ensure_ascii=False, separators=(",", ":"), sort_keys=True)
OUTPUT.write_text(
    "// Gerado a partir da planilha original da biblioteca.\n"
    f"window.BIBLIOTECA_CATALOG_METADATA=Object.freeze({payload});\n",
    encoding="utf-8",
)
print(f"{len(metadata)} títulos de referência gravados em {OUTPUT.name}")
